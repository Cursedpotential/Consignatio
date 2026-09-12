#!/usr/bin/env python3
"""
make_labeling_workbook.py — build the self-classification (ground-truth) Excel
workbook from the labeling CSV + category list.

Byline: Claude Code · Fable 5 · 2026-07-03

Owner self-classifies the pilot conversation to build a GOLD-STANDARD labeled set
(the first detection pass over-flagged; ground truth drives refinement + training).
Message-centric: every message gets a your_label dropdown so MISSES (false negatives)
are captured, not just review of AI guesses.

Inputs  (D:/casebible/exports/):
  labeling-workbook-imessage-2989.csv   — 1918 msgs + current AI flags + blank label cols
  _cats_for_dropdown.txt                — 164 'polarity:category_id' options
Output:
  labeling-workbook-imessage-2989.xlsx  — frozen header, autofilter, dropdowns:
     your_label   = NONE/UNSURE + the 164 categories
     your_severity= 0..10
     AI-flagged rows tinted so the over-flagging is visible while labeling.
Idempotent; re-run safe.
"""
import csv
from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

EXP = r"D:/casebible/exports"
SRC = f"{EXP}/labeling-workbook-imessage-2989.csv"
CATS = f"{EXP}/_cats_for_dropdown.txt"
OUT = f"{EXP}/labeling-workbook-imessage-2989.xlsx"

rows = list(csv.DictReader(open(SRC, encoding="utf-8")))
cats = [c.strip() for c in open(CATS, encoding="utf-8") if c.strip()]
label_opts = ["NONE (clean)", "UNSURE"] + cats     # what the owner picks per message

wb = Workbook()
ws = wb.active
ws.title = "label"

headers = ["seq", "when_utc", "who", "message", "ai_flagged", "ai_flag_count",
           "ai_max_sev", "your_label", "your_severity", "your_notes", "message_id"]
ws.append(headers)
for c in range(1, len(headers) + 1):
    cell = ws.cell(1, c)
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="2D3440")
    cell.alignment = Alignment(vertical="center")

me_fill    = PatternFill("solid", fgColor="E7F0FE")   # ME rows
flag_fill  = PatternFill("solid", fgColor="FDECEA")   # AI-flagged rows (see the over-flagging)
for r in rows:
    ws.append([r["seq"], r["when_utc"], r["who"], r["message"], r["ai_flagged"],
               int(r["ai_flag_count"] or 0), r["ai_max_sev"], "", "", "", r["message_id"]])
    i = ws.max_row
    if int(r["ai_flag_count"] or 0) > 0:
        for c in range(1, 8):
            ws.cell(i, c).fill = flag_fill
    elif r["who"] == "ME":
        ws.cell(i, 3).fill = me_fill

# dropdowns
n = ws.max_row
dv_label = DataValidation(type="list", formula1='"%s"' % ",".join(
    # Excel inline list caps ~255 chars; 164 cats blow that — use a helper column range instead.
    []), allow_blank=True)
# put option lists on a hidden sheet and reference ranges (handles long lists)
opt = wb.create_sheet("opts"); opt.sheet_state = "hidden"
for idx, v in enumerate(label_opts, 1):
    opt.cell(idx, 1, v)
for idx, v in enumerate([str(x) for x in range(0, 11)], 1):
    opt.cell(idx, 2, v)
label_ref = f"opts!$A$1:$A${len(label_opts)}"
sev_ref   = f"opts!$B$1:$B$11"

dvL = DataValidation(type="list", formula1=f"={label_ref}", allow_blank=True)
dvL.error = "Pick a category from the list (or NONE/UNSURE)"; dvL.errorTitle = "your_label"
dvS = DataValidation(type="list", formula1=f"={sev_ref}", allow_blank=True)
ws.add_data_validation(dvL); ws.add_data_validation(dvS)
dvL.add(f"H2:H{n}")   # your_label
dvS.add(f"I2:I{n}")   # your_severity

# widths + wrap + freeze + filter
widths = {"A":6,"B":15,"C":6,"D":70,"E":26,"F":6,"G":7,"H":22,"I":8,"J":30,"K":38}
for col, w in widths.items():
    ws.column_dimensions[col].width = w
for i in range(2, n + 1):
    ws.cell(i, 4).alignment = Alignment(wrap_text=True, vertical="top")
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:K{n}"

# instructions sheet
info = wb.create_sheet("READ ME", 0)
for i, line in enumerate([
    "SELF-CLASSIFICATION WORKBOOK — pilot conversation imessage:+18108532989 (1,918 messages)",
    "",
    "GOAL: your own read of the conversation = the ground truth. The machine's first pass over-flagged",
    "(374/1918 msgs, dominated by hedge_words/certainty_absolutes noise). We refine patterns against YOUR labels.",
    "",
    "HOW:",
    "  • Go to the 'label' tab. Red rows = the machine flagged them (ai_flagged shows its guess).",
    "  • For EVERY message, set 'your_label' from the dropdown:",
    "       NONE (clean)  = no notable behavior",
    "       UNSURE        = flag for discussion",
    "       <category>    = the behavior YOU see (pick the closest; label misses the machine skipped too)",
    "  • 'your_severity' 0-10 optional. 'your_notes' free text optional.",
    "  • You do NOT have to agree with ai_flagged — that's the point. Overrule it freely.",
    "",
    "WHY message-by-message (not just reviewing red rows): labeling the WHOLE thing captures what the",
    "machine MISSED (false negatives), which is half of what makes patterns better.",
    "",
    "When done: send it back. We load it as analysis.human_label (ground truth), score the 512 patterns",
    "against it (precision/recall per category), kill the noisy ones, and it becomes small-model training data.",
], 1):
    info.cell(i, 1, line)
info.column_dimensions["A"].width = 110

wb.save(OUT)
print("wrote", OUT, "|", n - 1, "message rows |", len(cats), "categories in dropdown")
