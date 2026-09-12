from __future__ import annotations

import csv
import io
import json
import re
import xml.etree.ElementTree as ET
from email import policy
from email.parser import BytesParser
from pathlib import Path

from bs4 import BeautifulSoup
from charset_normalizer import from_bytes
from docx import Document
from pypdf import PdfReader
from striprtf.striprtf import rtf_to_text

from .models import ExtractedText

PLAIN_TEXT_EXTENSIONS = {
    ".csv",
    ".jsonl",
    ".log",
    ".markdown",
    ".md",
    ".rst",
    ".text",
    ".tsv",
    ".txt",
    ".yaml",
    ".yml",
}


def _normalize(text: str) -> str:
    text = text.replace("\x00", "")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{4,}", "\n\n\n", text)
    return text.strip()


def _decode(content: bytes) -> str:
    best = from_bytes(content).best()
    if best is None:
        return content.decode("utf-8", errors="replace")
    return str(best)


def _html_to_text(html: str) -> str:
    soup = BeautifulSoup(html, "html.parser")
    for element in soup(["script", "style", "noscript"]):
        element.decompose()
    return soup.get_text("\n", strip=True)


def _extract_email(content: bytes) -> str:
    message = BytesParser(policy=policy.default).parsebytes(content)
    headers = []
    for name in ("Date", "From", "To", "Cc", "Subject"):
        if message.get(name):
            headers.append(f"{name}: {message.get(name)}")

    bodies: list[str] = []
    parts = message.walk() if message.is_multipart() else [message]
    for part in parts:
        if part.get_content_disposition() == "attachment":
            continue
        content_type = part.get_content_type()
        if content_type not in {"text/plain", "text/html"}:
            continue
        try:
            body = part.get_content()
        except (LookupError, UnicodeError):
            payload = part.get_payload(decode=True) or b""
            body = _decode(payload)
        if content_type == "text/html":
            body = _html_to_text(str(body))
        if str(body).strip():
            bodies.append(str(body).strip())
    return "\n".join(headers + [""] + bodies)


def _extract_delimited(content: bytes, delimiter: str) -> str:
    text = _decode(content)
    rows = csv.reader(io.StringIO(text), delimiter=delimiter)
    return "\n".join(" | ".join(cell.strip() for cell in row) for row in rows)


def extract_text(path: Path, content: bytes) -> ExtractedText:
    """Extract text without changing the source or invoking OCR/vision."""

    extension = path.suffix.casefold()
    media_type = {
        ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ".eml": "message/rfc822",
        ".html": "text/html",
        ".htm": "text/html",
        ".json": "application/json",
        ".pdf": "application/pdf",
        ".rtf": "application/rtf",
        ".xml": "application/xml",
    }.get(extension, "text/plain")

    try:
        page_count: int | None = None
        method = "decoded_text"
        notes: tuple[str, ...] = ()

        if extension == ".pdf":
            reader = PdfReader(io.BytesIO(content))
            page_count = len(reader.pages)
            pages = [(page.extract_text() or "").strip() for page in reader.pages]
            text = "\n\n".join(page for page in pages if page)
            method = "pypdf_text_layer"
            if len(text.strip()) < 40:
                return ExtractedText(
                    text="",
                    media_type=media_type,
                    page_count=page_count,
                    extraction_method=method,
                    status="skipped_no_text",
                    notes=("PDF has no usable text layer; OCR is intentionally deferred.",),
                )
        elif extension == ".docx":
            document = Document(io.BytesIO(content))
            blocks = [p.text for p in document.paragraphs]
            for table in document.tables:
                blocks.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
            text = "\n".join(blocks)
            method = "python_docx"
        elif extension in {".html", ".htm"}:
            text = _html_to_text(_decode(content))
            method = "beautifulsoup_text"
        elif extension == ".eml":
            text = _extract_email(content)
            method = "email_parser"
        elif extension == ".rtf":
            text = rtf_to_text(_decode(content), errors="ignore")
            method = "striprtf"
        elif extension == ".json":
            decoded = _decode(content)
            try:
                text = json.dumps(json.loads(decoded), ensure_ascii=False, indent=2)
                method = "json_normalized"
            except json.JSONDecodeError:
                text = decoded
                notes = ("Invalid JSON indexed as decoded text.",)
        elif extension == ".xml":
            decoded = _decode(content)
            try:
                root = ET.fromstring(decoded)
                text = "\n".join(value.strip() for value in root.itertext() if value.strip())
                method = "xml_itertext"
            except ET.ParseError:
                text = decoded
                notes = ("Invalid XML indexed as decoded text.",)
        elif extension == ".csv":
            text = _extract_delimited(content, ",")
            method = "csv_rows"
        elif extension == ".tsv":
            text = _extract_delimited(content, "\t")
            method = "tsv_rows"
        elif extension in PLAIN_TEXT_EXTENSIONS:
            text = _decode(content)
        else:
            return ExtractedText(
                text="",
                media_type=media_type,
                page_count=None,
                extraction_method="unsupported",
                status="unsupported",
                notes=(f"Unsupported text extension: {extension or '<none>'}",),
            )

        text = _normalize(text)
        if not text:
            return ExtractedText(
                text="",
                media_type=media_type,
                page_count=page_count,
                extraction_method=method,
                status="skipped_no_text",
                notes=notes + ("No usable text was extracted.",),
            )
        return ExtractedText(
            text=text,
            media_type=media_type,
            page_count=page_count,
            extraction_method=method,
            status="indexed",
            notes=notes,
        )
    except Exception as exc:
        raise ValueError(f"Failed to extract {path.name}: {type(exc).__name__}: {exc}") from exc
